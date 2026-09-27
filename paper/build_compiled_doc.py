from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "AQUA_Journal_Paper_Package.docx"
SOURCES = [
    ("Manuscript Draft", "manuscript.md"),
    ("Project and Experiment Audit", "01_project_and_experiment_audit.md"),
    ("Literature Review", "02_literature_review.md"),
    ("Citation Audit", "03_citation_audit.md"),
    ("IEEE Pre Submission Review", "04_ieee_pre_submission_review.md"),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def add_table(doc, rows):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            cell = cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if row_index == 0:
                set_cell_shading(cell, "1F4E79")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F7FA")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(value)
            run.font.size = Pt(9)
            if row_index == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph()


def parse_markdown(doc, text):
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("|") and "|" in line[1:]:
            block = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                current = lines[index].strip()
                if not set(current.replace("|", "").replace("-", "").replace(":", "").strip()):
                    index += 1
                    continue
                block.append([part.strip() for part in current.strip("|").split("|")])
                index += 1
            if block and len({len(row) for row in block}) == 1:
                add_table(doc, block)
            else:
                for row in block:
                    doc.add_paragraph(" | ".join(row))
            continue
        if line.startswith("#"):
            marks = len(line) - len(line.lstrip("#"))
            heading = line[marks:].strip()
            style = "Heading 1" if marks == 1 else "Heading 2" if marks == 2 else "Heading 3"
            paragraph = doc.add_paragraph(heading, style=style)
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            index += 1
            continue
        if line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
            index += 1
            continue
        if line[:3].isdigit() and ". " in line[:5]:
            doc.add_paragraph(line.split(". ", 1)[1], style="List Number")
            index += 1
            continue
        if line.startswith("**") and line.endswith("**"):
            paragraph = doc.add_paragraph()
            paragraph.add_run(line.strip("*")).bold = True
            index += 1
            continue
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.12
        paragraph.add_run(line.replace("**", "").replace("`", ""))
        index += 1


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph("AQUA Journal Paper Package", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Evidence bounded manuscript draft and supporting audits")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    doc.add_paragraph("This document consolidates the current manuscript, project and experiment audit, targeted literature review, citation audit, and IEEE style pre submission review. It is a research package, not a submission ready article: required experiments and venue specific fields remain explicitly marked.")
    doc.add_paragraph("Prepared from the audited AQUA repository on 24 September 2026.")

    for section_title, filename in SOURCES:
        doc.add_page_break()
        doc.add_paragraph(section_title, style="Heading 1")
        parse_markdown(doc, (ROOT / filename).read_text(encoding="utf-8"))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("AQUA journal paper package")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
